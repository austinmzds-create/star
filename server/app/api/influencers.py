"""达人库:智能录入(粘贴解析)、档案、定级(留痕)、商务归属。"""
from io import BytesIO
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..db import get_db
from ..deps import current_user, owns_or_admin
from ..models import (AccessGrant, Cooperation, Influencer, LevelChangeLog,
                      OperationLog, OrderRecord, Product, Promotion,
                      SampleOrder, User, VideoTask)
from ..services import levels, storage
from ..services.oplog import log_op
from ..services.parser import parse_influencer_text
from ..services.sample_orders import dedupe_sample_rows
from ..services.tracking import refresh_if_needed

router = APIRouter(prefix="/api/influencers", tags=["influencers"])

TRACKED_FIELDS = ("level", "commission_tier", "promo_mode", "owner_bd_id")
IDENTITY_FIELDS = ("douyin_uid", "douyin_id", "phone", "cooperation_code")
VALID_LEVELS = ("L1", "L2", "L3")

IMPORT_COLUMNS = [
    ("昵称", "nickname", True),
    ("抖音号", "douyin_id", True),
    ("UID", "douyin_uid", False),
    ("主页链接", "homepage_url", False),
    ("手机号", "phone", False),
    ("粉丝数", "fans_count", False),
    ("近30天GMV", "gmv_30d", False),
    ("内容品类(逗号分隔)", "category_tags", False),
    ("拍摄类型", "shoot_type", False),
    ("合作码", "cooperation_code", False),
    ("收件人", "real_name", False),
    ("收件地址", "default_address", False),
    ("等级(L1/L2/L3)", "level", False),
    ("标签(逗号分隔)", "tags", False),
    ("数据来源", "data_source", False),
    ("来源备注", "source_note", False),
    ("归属商务ID", "owner_bd_id", False),
    ("归属商务手机号", "owner_bd_phone", False),
]


def scope(db_query, user: User):
    """商务只见自己的达人;管理员全量"""
    if user.role == "admin":
        return db_query
    return db_query.where(Influencer.owner_bd_id == user.id)


class ParseIn(BaseModel):
    text: str


# 可在「重复确认」里对比/更新的档案字段(排除归属、管理员备注等敏感/管控字段)
DIFF_FIELDS = [
    ("nickname", "昵称"), ("douyin_id", "抖音号"), ("douyin_uid", "UID"),
    ("phone", "手机号"), ("real_name", "收件人"), ("default_address", "收件地址"),
    ("homepage_url", "主页"), ("fans_count", "粉丝数"), ("gmv_30d", "近30天GMV"),
    ("shoot_type", "拍摄类型"), ("cooperation_code", "合作码"),
    ("data_source", "数据来源"), ("category_tags", "内容品类"),
]


def _norm_for_compare(field, value):
    """把 None/空串/列表统一成可比较的展示值。"""
    if field == "category_tags":
        if isinstance(value, list):
            return "、".join(str(v) for v in value)
        return (value or "").strip() if isinstance(value, str) else (value or "")
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value)


def _build_diff(existing: Influencer, incoming: dict) -> list[dict]:
    """逐字段对比「已有档案」与「本次录入」,只返回本次有值且与已有不同的字段。"""
    diff = []
    for field, label in DIFF_FIELDS:
        new_raw = incoming.get(field)
        if new_raw is None or (isinstance(new_raw, str) and not new_raw.strip()):
            continue
        old_show = _norm_for_compare(field, getattr(existing, field, None))
        new_show = _norm_for_compare(field, new_raw)
        if old_show != new_show:
            diff.append({"field": field, "label": label, "old": old_show, "new": new_show})
    return diff


def _identity_conditions(fields: dict, exclude_id: int | None = None):
    conds = []
    for field in IDENTITY_FIELDS:
        raw = fields.get(field)
        value = normalize_douyin(raw) if field == "douyin_id" else _clean_identity(raw)
        if value:
            conds.append(getattr(Influencer, field) == value)
    stmt = select(Influencer)
    if conds:
        stmt = stmt.where(or_(*conds))
        if exclude_id is not None:
            stmt = stmt.where(Influencer.id != exclude_id)
    return stmt if conds else None


def _find_duplicate(db: Session, fields: dict, exclude_id: int | None = None) -> Influencer | None:
    stmt = _identity_conditions(fields, exclude_id)
    if stmt is None:
        return None
    return db.scalars(stmt.order_by(Influencer.archived, Influencer.id)).first()


def _clean_identity(value):
    if isinstance(value, str):
        value = value.strip()
    return value or None


def normalize_douyin(value):
    """抖音号规范化:去首尾空格 + 去前导 @(方案B 需求2 唯一标识口径)。"""
    v = _clean_identity(value)
    if not v:
        return None
    return v.lstrip("@").strip() or None


def _normalize_create(body) -> dict:
    data = body.model_dump()
    data["nickname"] = (data.get("nickname") or "").strip()
    data["douyin_id"] = normalize_douyin(data.get("douyin_id"))
    data["douyin_uid"] = _clean_identity(data.get("douyin_uid"))
    data["phone"] = _clean_identity(data.get("phone"))
    data["cooperation_code"] = _clean_identity(data.get("cooperation_code"))
    data["data_source"] = _clean_identity(data.get("data_source"))
    data["source_note"] = _clean_identity(data.get("source_note"))
    data["level"] = (data.get("level") or "L1").strip().upper()
    if not data["nickname"]:
        raise HTTPException(400, "达人昵称不能为空")
    if not data["douyin_id"]:
        raise HTTPException(400, "抖音号必填,它是达人唯一识别标识")
    if data["level"] not in VALID_LEVELS:
        raise HTTPException(400, "等级只能填写 L1/L2/L3")
    return data


def _cell_to_str(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    return text or None


def _split_list(value) -> list[str] | None:
    text = _cell_to_str(value)
    if not text:
        return None
    import re
    parts = [p.strip() for p in re.split(r"[,，;；、|\n]+", text) if p.strip()]
    return parts or None


def _parse_int(value, field_name: str) -> int | None:
    text = _cell_to_str(value)
    if not text:
        return None
    text = text.replace(",", "").replace("，", "").replace(" ", "")
    multiplier = Decimal("1")
    if text[-1:].lower() == "w" or text.endswith("万"):
        multiplier = Decimal("10000")
        text = text[:-1]
    elif text.endswith("亿"):
        multiplier = Decimal("100000000")
        text = text[:-1]
    text = text.replace("¥", "").replace("￥", "").replace("元", "")
    try:
        value_decimal = Decimal(text) * multiplier
    except Exception as exc:
        raise ValueError(f"{field_name}格式不正确") from exc
    if value_decimal < 0:
        raise ValueError(f"{field_name}不能为负数")
    return int(value_decimal)


def _normal_header(value) -> str:
    text = _cell_to_str(value) or ""
    return (text.replace("（", "(").replace("）", ")")
            .replace("*", "").replace(" ", "").strip())


def _import_header_aliases() -> dict[str, str]:
    aliases = {}
    for label, key, _required in IMPORT_COLUMNS:
        aliases[_normal_header(label)] = key
        aliases[_normal_header(key)] = key
    aliases.update({
        "昵称(必填)": "nickname",
        "达人昵称": "nickname",
        "抖音号(必填唯一)": "douyin_id",
        "抖音号(必填)": "douyin_id",
        "抖音ID": "douyin_id",
        "达人UID": "douyin_uid",
        "主页": "homepage_url",
        "主页URL": "homepage_url",
        "手机号/联系电话": "phone",
        "联系电话": "phone",
        "粉丝": "fans_count",
        "GMV": "gmv_30d",
        "30天GMV": "gmv_30d",
        "品类": "category_tags",
        "内容品类": "category_tags",
        "达人类型": "shoot_type",
        "收货人": "real_name",
        "收货地址": "default_address",
        "等级": "level",
        "标签": "tags",
        "来源": "data_source",
        "官方来源": "data_source",
        "数据来源/官方来源": "data_source",
        "备注": "source_note",
        "官方备注": "source_note",
        "归属商务": "owner_bd_phone",
        "商务手机号": "owner_bd_phone",
    })
    return {_normal_header(k): v for k, v in aliases.items()}


def _owner_for_import(db: Session, user: User, row_data: dict) -> int | None:
    if user.role != "admin":
        return user.id
    owner_id = _parse_int(row_data.get("owner_bd_id"), "归属商务ID")
    owner_phone = _clean_identity(_cell_to_str(row_data.get("owner_bd_phone")))
    if owner_id:
        owner = db.get(User, owner_id)
        if not owner or not owner.is_active:
            raise ValueError("归属商务ID不存在或已停用")
        return owner.id
    if owner_phone:
        owner = db.scalars(select(User).where(User.phone == owner_phone, User.is_active.is_(True))).first()
        if not owner:
            raise ValueError("归属商务手机号不存在或已停用")
        return owner.id
    return user.id


def _payload_from_import_row(db: Session, user: User, row_data: dict) -> tuple[dict, list[str] | None, int | None]:
    nickname = _cell_to_str(row_data.get("nickname"))
    douyin_id = normalize_douyin(_cell_to_str(row_data.get("douyin_id")))
    if not nickname:
        raise ValueError("昵称必填")
    if not douyin_id:
        raise ValueError("抖音号必填,它是达人唯一识别标识")
    level = (_cell_to_str(row_data.get("level")) or "L1").upper()
    if level not in VALID_LEVELS:
        raise ValueError("等级只能填写 L1/L2/L3")

    payload = {
        "nickname": nickname,
        "douyin_id": douyin_id,
        "douyin_uid": _clean_identity(_cell_to_str(row_data.get("douyin_uid"))),
        "homepage_url": _cell_to_str(row_data.get("homepage_url")),
        "real_name": _cell_to_str(row_data.get("real_name")),
        "phone": _clean_identity(_cell_to_str(row_data.get("phone"))),
        "fans_count": _parse_int(row_data.get("fans_count"), "粉丝数"),
        "gmv_30d": _parse_int(row_data.get("gmv_30d"), "近30天GMV"),
        "category_tags": _split_list(row_data.get("category_tags")),
        "shoot_type": _cell_to_str(row_data.get("shoot_type")),
        "cooperation_code": _clean_identity(_cell_to_str(row_data.get("cooperation_code"))),
        "default_address": _cell_to_str(row_data.get("default_address")),
        "homepage_raw": None,
        "data_source": _cell_to_str(row_data.get("data_source")),
        "source_note": _cell_to_str(row_data.get("source_note")),
        "level": level,
        "source": "import",
    }
    tags = _split_list(row_data.get("tags"))
    owner_bd_id = _owner_for_import(db, user, row_data)
    return payload, tags, owner_bd_id


def _create_imported_influencer(db: Session, user: User, payload: dict,
                                tags: list[str] | None, owner_bd_id: int | None) -> Influencer:
    existing = _find_duplicate(db, payload)
    if existing:
        if owns_or_admin(user, existing.owner_bd_id):
            raise ValueError(f"达人已存在(ID {existing.id}),请打开已有档案维护")
        owner = db.get(User, existing.owner_bd_id) if existing.owner_bd_id else None
        raise ValueError(f"该达人已由{owner.display_name if owner else '其他商务'}对接,请勿重复导入")
    cfg = levels.effective_config(db, payload["level"])
    inf = Influencer(**payload,
                     commission_tier=cfg.commission_tier if cfg else Decimal("5"),
                     owner_bd_id=owner_bd_id,
                     tags=tags)
    db.add(inf)
    db.flush()
    db.add(Cooperation(
        influencer_id=inf.id,
        round_no=1,
        level_snapshot=inf.level,
        commission_tier_snapshot=inf.commission_tier,
        promo_mode_snapshot=inf.promo_mode,
    ))
    db.commit()
    return inf


@router.post("/parse")
async def parse(body: ParseIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """粘贴自我介绍 → 解析字段 + 撞库查重(老达人提示第N次合作)"""
    result = await parse_influencer_text(body.text)
    f = result["fields"]
    dup = None
    existing = _find_duplicate(db, f)
    if existing:
        if owns_or_admin(user, existing.owner_bd_id):
            dup = {"id": existing.id, "nickname": existing.nickname,
                   "round_count": len(existing.cooperations),
                   "owner_bd_id": existing.owner_bd_id,
                   "diff": _build_diff(existing, f)}
        else:
            # 已被其他商务对接:只提示归属,不泄漏其达人档案(数据隔离)
            owner = db.get(User, existing.owner_bd_id) if existing.owner_bd_id else None
            dup = {"owned_by_other_bd": True,
                   "owner_bd_name": owner.display_name if owner else None}
    return {**result, "duplicate": dup}


class CreateIn(BaseModel):
    nickname: str
    douyin_id: str | None = None
    douyin_uid: str | None = None
    homepage_url: str | None = None
    real_name: str | None = None
    phone: str | None = None
    fans_count: int | None = None
    gmv_30d: int | None = None
    category_tags: list[str] | None = None
    shoot_type: str | None = None
    raw_intro: str | None = None
    cooperation_code: str | None = None
    default_address: str | None = None
    homepage_raw: str | None = None
    data_source: str | None = None
    source_note: str | None = None
    level: str = "L1"
    source: str = "bd"


@router.post("")
def create(body: CreateIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    data = _normalize_create(body)
    existing = _find_duplicate(db, data)
    if existing:
        if owns_or_admin(user, existing.owner_bd_id):
            raise HTTPException(409, f"达人已存在(ID {existing.id}),请打开已有档案继续合作")
        owner = db.get(User, existing.owner_bd_id) if existing.owner_bd_id else None
        raise HTTPException(409, f"该达人已由{owner.display_name if owner else '其他商务'}对接,请勿重复建档")
    cfg = levels.effective_config(db, body.level)
    inf = Influencer(**data,
                     commission_tier=cfg.commission_tier if cfg else Decimal("5"),
                     owner_bd_id=user.id)
    db.add(inf)
    db.commit()
    levels.new_cooperation(db, inf)  # 录入即开第1轮合作(写入快照)
    return {"id": inf.id}


@router.get("/import-template")
def import_template(user: User = Depends(current_user)):
    """下载达人批量导入 Excel 模板。"""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "达人导入"
    headers = [label for label, _key, _required in IMPORT_COLUMNS]
    ws.append(headers)
    ws.append([
        "二宝妈妈", "douyin_123456", "1122334455", "https://www.douyin.com/user/xxx",
        "15000000000", "1.2万", "23000", "母婴,儿童", "口播", "BY123",
        "张三", "上海市浦东新区示例路 1 号", "L1", "高潜,母婴",
        "官方达人库", "2026-07-13 官方后台导出,待商务复核", "", "",
    ])
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    for idx, (_label, _key, required) in enumerate(IMPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        if required:
            cell.fill = required_fill
            cell.comment = Comment("必填。抖音号是唯一识别标识,重复会导入失败。", "system")
        ws.column_dimensions[cell.column_letter].width = 18
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"L1,L2,L3"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("M2:M2000")
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    headers = {"Content-Disposition": "attachment; filename=influencer_import_template.xlsx"}
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.post("/import")
async def import_influencers(file: UploadFile = File(...),
                             user: User = Depends(current_user),
                             db: Session = Depends(get_db)):
    """批量导入达人:逐行校验,成功行落库,失败行回传。"""
    from openpyxl import load_workbook

    filename = file.filename or ""
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式的 Excel 文件")
    content = await file.read()
    if not content:
        raise HTTPException(400, "上传文件为空")
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "Excel 文件不能超过 5MB")
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, "Excel 文件解析失败,请使用系统模板重新导入") from exc
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise HTTPException(400, "Excel 为空") from exc
    aliases = _import_header_aliases()
    col_map = {}
    for idx, header in enumerate(header_row):
        key = aliases.get(_normal_header(header))
        if key:
            col_map[idx] = key
    required = {key for _label, key, is_required in IMPORT_COLUMNS if is_required}
    missing = required - set(col_map.values())
    if missing:
        labels = {key: label for label, key, _required in IMPORT_COLUMNS}
        raise HTTPException(400, "模板缺少必填列:" + ",".join(labels.get(k, k) for k in missing))

    created = []
    failures = []
    seen_douyin = set()
    total = 0
    for row_no, row in enumerate(rows, start=2):
        if row_no > 2001:
            failures.append({"row": row_no, "reason": "单次最多导入 2000 行,剩余行请拆分文件"})
            break
        if not any(_cell_to_str(value) for value in row):
            continue
        total += 1
        row_data = {key: row[idx] if idx < len(row) else None for idx, key in col_map.items()}
        try:
            payload, tags, owner_bd_id = _payload_from_import_row(db, user, row_data)
            douyin_key = payload["douyin_id"].lower()
            if douyin_key in seen_douyin:
                raise ValueError("Excel 内抖音号重复")
            inf = _create_imported_influencer(db, user, payload, tags, owner_bd_id)
            seen_douyin.add(douyin_key)
            created.append({
                "row": row_no,
                "id": inf.id,
                "nickname": inf.nickname,
                "douyin_id": inf.douyin_id,
            })
        except ValueError as exc:
            db.rollback()
            failures.append({
                "row": row_no,
                "nickname": _cell_to_str(row_data.get("nickname")),
                "douyin_id": _cell_to_str(row_data.get("douyin_id")),
                "reason": str(exc),
            })
        except Exception:
            db.rollback()
            failures.append({
                "row": row_no,
                "nickname": _cell_to_str(row_data.get("nickname")),
                "douyin_id": _cell_to_str(row_data.get("douyin_id")),
                "reason": "系统处理失败,请检查该行数据后重试",
            })
    return {
        "total": total,
        "success_count": len(created),
        "failed_count": len(failures),
        "created": created,
        "failures": failures,
    }


@router.get("")
def list_influencers(q: str | None = None, level: str | None = None,
                     commission_tier: float | None = None,
                     owner_bd_id: int | None = None, tag: str | None = None,
                     include_archived: bool = False,
                     page: int = 1, page_size: int = 50,
                     user: User = Depends(current_user), db: Session = Depends(get_db)):
    from sqlalchemy import func
    base = scope(select(Influencer), user)
    if not include_archived:
        base = base.where(Influencer.archived.is_not(True))
    if q:
        like = f"%{q}%"
        base = base.where(or_(Influencer.nickname.like(like), Influencer.douyin_id.like(like),
                              Influencer.douyin_uid.like(like), Influencer.phone.like(like),
                              Influencer.cooperation_code.like(like),
                              Influencer.data_source.like(like)))
    if level:
        base = base.where(Influencer.level == level)
    if commission_tier is not None:
        base = base.where(Influencer.commission_tier == Decimal(str(commission_tier)))
    if owner_bd_id is not None:
        base = base.where(Influencer.owner_bd_id == owner_bd_id)
    total = db.scalar(select(func.count()).select_from(base.subquery()))
    page = max(1, page)
    page_size = min(max(1, page_size), 200)
    rows = db.scalars(base.order_by(Influencer.updated_at.desc())
                      .offset((page - 1) * page_size).limit(page_size)).all()
    # 归属商务名字(一次性查出映射)
    bd_ids = {r.owner_bd_id for r in rows if r.owner_bd_id}
    names = {u.id: u.display_name for u in
             db.scalars(select(User).where(User.id.in_(bd_ids or [0]))).all()}
    items = [{"id": r.id, "nickname": r.nickname, "douyin_id": r.douyin_id,
              "fans_count": r.fans_count, "gmv_30d": r.gmv_30d, "level": r.level,
              "commission_tier": float(r.commission_tier), "promo_mode": r.promo_mode,
              "tags": r.tags, "round_count": len(r.cooperations),
              "owner_bd_id": r.owner_bd_id, "owner_bd_name": names.get(r.owner_bd_id),
              "source": r.source, "data_source": r.data_source,
              "updated_at": r.updated_at.isoformat()}
             for r in rows]
    # tag 过滤(tags 存 JSON,DB 层不易过滤,内存过滤本页)
    if tag:
        items = [i for i in items if i["tags"] and tag in i["tags"]]
    return {"items": items, "total": total, "page": page, "page_size": page_size}


class UpdateIn(BaseModel):
    level: str | None = None
    commission_tier: float | None = None
    promo_mode: str | None = None
    owner_bd_id: int | None = None  # 转移分配:仅管理员
    tags: list[str] | None = None
    gmv_30d: int | None = None
    shoot_type: str | None = None
    reason: str | None = None       # 调级/调档原因(留痕)
    archived: bool | None = None    # 停用/启用
    # 核心档案字段(录错可改)
    nickname: str | None = None
    douyin_id: str | None = None
    douyin_uid: str | None = None
    real_name: str | None = None
    phone: str | None = None
    fans_count: int | None = None
    category_tags: list[str] | None = None
    cooperation_code: str | None = None
    default_address: str | None = None
    homepage_url: str | None = None
    data_source: str | None = None
    source_note: str | None = None
    admin_note: str | None = None       # 管理员私密备注:仅管理员可看可改


CORE_FIELDS = ("nickname", "douyin_id", "douyin_uid", "real_name", "phone",
               "fans_count", "category_tags", "cooperation_code",
               "default_address", "homepage_url", "data_source",
               "source_note", "archived")

FIELD_LABELS = {
    "level": "等级", "commission_tier": "佣金档", "promo_mode": "投流方式",
    "owner_bd_id": "归属商务", "nickname": "昵称", "douyin_id": "抖音号",
    "douyin_uid": "UID", "real_name": "收件人", "phone": "手机号",
    "fans_count": "粉丝数", "category_tags": "内容品类", "cooperation_code": "合作码",
    "default_address": "收件地址", "homepage_url": "主页", "data_source": "数据来源",
    "source_note": "来源备注", "archived": "启用状态", "tags": "标签",
    "gmv_30d": "近30天GMV", "shoot_type": "拍摄类型",
}


def _display(field, value):
    if field == "archived":
        return "停用" if value else "启用"
    if value is None or value == "":
        return "空"
    return str(value)


@router.patch("/{influencer_id}")
def update(influencer_id: int, body: UpdateIn,
           user: User = Depends(current_user), db: Session = Depends(get_db)):
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    if body.owner_bd_id is not None and user.role != "admin":
        raise HTTPException(403, "无权转移达人(仅管理员可分配归属)")
    if "admin_note" in body.model_fields_set and user.role != "admin":
        raise HTTPException(403, "管理员备注仅管理员可维护")

    for field in TRACKED_FIELDS:
        new_val = getattr(body, field, None)
        if new_val is None:
            continue
        old_val = getattr(inf, field)
        if str(old_val) != str(new_val):
            db.add(LevelChangeLog(influencer_id=inf.id, field=field,
                                  old_value=str(old_val), new_value=str(new_val),
                                  reason=body.reason, changed_by=user.id))
            setattr(inf, field, new_val if field != "commission_tier" else Decimal(str(new_val)))
            # 统一时间轴留痕:佣金/归属为业务关键事件,其余定级/投流方式记为档案变更
            label = FIELD_LABELS.get(field, field)
            event = {"commission_tier": "commission_changed",
                     "owner_bd_id": "owner_transferred"}.get(field, "profile_changed")
            log_op(db, influencer_id=inf.id, event_type=event, actor=user,
                   summary=f"{user.display_name} 修改{label}:{_display(field, old_val)}→{_display(field, new_val)}"
                           + (f"(原因:{body.reason})" if body.reason else ""),
                   detail={"field": field, "old": str(old_val), "new": str(new_val),
                           "reason": body.reason})
        # 调级时联动默认佣金档(可再被单独覆盖)
        if field == "level" and body.commission_tier is None:
            cfg = levels.effective_config(db, str(new_val))
            if cfg:
                inf.commission_tier = cfg.commission_tier
    for field in ("tags", "gmv_30d", "shoot_type"):
        if getattr(body, field) is not None:
            setattr(inf, field, getattr(body, field))
    # 核心档案字段(录错可改);身份字段统一撞库校验,避免重复达人。
    identity_updates = {}
    for field in IDENTITY_FIELDS:
        if getattr(body, field, None) is not None:
            raw = getattr(body, field)
            identity_updates[field] = normalize_douyin(raw) if field == "douyin_id" else _clean_identity(raw)
    if "douyin_id" in identity_updates and not identity_updates["douyin_id"]:
        raise HTTPException(400, "抖音号不能为空")
    if identity_updates:
        clash = _find_duplicate(db, identity_updates, exclude_id=inf.id)
        if clash and clash.archived is not True:
            raise HTTPException(400, f"该达人已存在(ID {clash.id}),不能重复使用相同抖音号/UID/手机号/合作码")
    for field in CORE_FIELDS:
        if getattr(body, field) is not None:
            value = getattr(body, field)
            if field == "douyin_id":
                value = normalize_douyin(value)
            elif field in IDENTITY_FIELDS:
                value = _clean_identity(value)
            old_val = getattr(inf, field)
            if str(old_val) != str(value):
                label = FIELD_LABELS.get(field, field)
                log_op(db, influencer_id=inf.id, event_type="profile_changed", actor=user,
                       summary=f"{user.display_name} 修改{label}:{_display(field, old_val)}→{_display(field, value)}",
                       detail={"field": field, "old": _display(field, old_val),
                               "new": _display(field, value)})
            setattr(inf, field, value)
    if "admin_note" in body.model_fields_set:
        inf.admin_note = body.admin_note
    db.commit()
    return {"ok": True}


class ApplyUpdateIn(BaseModel):
    fields: dict
    reason: str | None = None


@router.post("/{influencer_id}/apply-update")
def apply_update(influencer_id: int, body: ApplyUpdateIn,
                 user: User = Depends(current_user), db: Session = Depends(get_db)):
    """重复确认:按用户勾选的字段更新已有达人,逐字段落 OperationLog(方案B 需求2)。"""
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    allowed = {f for f, _ in DIFF_FIELDS}
    updates = {k: v for k, v in (body.fields or {}).items() if k in allowed}
    if not updates:
        return {"ok": True, "changed": 0}
    # 身份字段规范化 + 唯一校验(避免撞库产生重复达人)
    identity_updates = {}
    for field in IDENTITY_FIELDS:
        if field in updates:
            raw = updates[field]
            val = normalize_douyin(raw) if field == "douyin_id" else _clean_identity(raw)
            updates[field] = val
            if val:
                identity_updates[field] = val
    if "douyin_id" in updates and not updates["douyin_id"]:
        raise HTTPException(400, "抖音号不能为空")
    if identity_updates:
        clash = _find_duplicate(db, identity_updates, exclude_id=inf.id)
        if clash and clash.archived is not True:
            raise HTTPException(400, f"该达人已存在(ID {clash.id}),不能重复使用相同抖音号/UID/手机号/合作码")
    changed = 0
    for field, value in updates.items():
        if field == "category_tags" and isinstance(value, str):
            import re
            value = [p.strip() for p in re.split(r"[,，;；、|\n]+", value) if p.strip()] or None
        old_val = getattr(inf, field, None)
        if _norm_for_compare(field, old_val) == _norm_for_compare(field, value):
            continue
        label = FIELD_LABELS.get(field, field)
        log_op(db, influencer_id=inf.id, event_type="profile_changed", actor=user,
               summary=f"{user.display_name} 修改{label}:{_norm_for_compare(field, old_val) or '空'}"
                       f"→{_norm_for_compare(field, value) or '空'}"
                       + (f"(原因:{body.reason})" if body.reason else ""),
               detail={"field": field, "old": _norm_for_compare(field, old_val),
                       "new": _norm_for_compare(field, value), "reason": body.reason})
        setattr(inf, field, value)
        changed += 1
    db.commit()
    return {"ok": True, "changed": changed}


@router.delete("/{influencer_id}")
def delete_influencer(influencer_id: int, user: User = Depends(current_user),
                      db: Session = Depends(get_db)):
    """硬删除:仅在无寄样/视频业务记录时允许(录入即开的空合作轮次会一并清理);
    已有寄样/视频请改用「停用」(archived) 以保留历史。"""
    from ..models import AccessGrant
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    coop_ids = [c.id for c in inf.cooperations] or [0]
    has_sample = db.scalar(select(SampleOrder.id)
                           .where(SampleOrder.cooperation_id.in_(coop_ids)).limit(1))
    has_video = db.scalar(select(VideoTask.id)
                          .where(VideoTask.cooperation_id.in_(coop_ids)).limit(1))
    if has_sample or has_video:
        raise HTTPException(400, "该达人已有寄样/视频记录,不能删除;请改用「停用」")
    db.query(AccessGrant).filter(AccessGrant.influencer_id == inf.id).delete()
    db.query(LevelChangeLog).filter(LevelChangeLog.influencer_id == inf.id).delete()
    for c in inf.cooperations:
        db.delete(c)
    db.delete(inf)
    db.commit()
    return {"ok": True}


@router.get("/{influencer_id}")
def detail(influencer_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    logs = db.scalars(select(LevelChangeLog).where(LevelChangeLog.influencer_id == inf.id)
                      .order_by(LevelChangeLog.changed_at.desc()).limit(50)).all()
    data = {
        "id": inf.id, "nickname": inf.nickname, "douyin_id": inf.douyin_id,
        "douyin_uid": inf.douyin_uid, "homepage_url": inf.homepage_url,
        "real_name": inf.real_name, "phone": inf.phone,
        "fans_count": inf.fans_count, "gmv_30d": inf.gmv_30d,
        "category_tags": inf.category_tags, "shoot_type": inf.shoot_type,
        "level": inf.level, "commission_tier": float(inf.commission_tier),
        "promo_mode": inf.promo_mode, "tags": inf.tags, "source": inf.source,
        "data_source": inf.data_source, "source_note": inf.source_note,
        "raw_intro": inf.raw_intro, "owner_bd_id": inf.owner_bd_id,
        "cooperation_code": inf.cooperation_code, "default_address": inf.default_address,
        "homepage_raw": inf.homepage_raw, "archived": inf.archived,
        "cooperations": [{"id": c.id, "round_no": c.round_no, "status": c.status,
                          "level_snapshot": c.level_snapshot,
                          "commission_tier_snapshot": float(c.commission_tier_snapshot),
                          "created_at": c.created_at.isoformat()} for c in inf.cooperations],
        "change_logs": [{"field": l.field, "old": l.old_value, "new": l.new_value,
                         "reason": l.reason, "at": l.changed_at.isoformat()} for l in logs],
    }
    if user.role == "admin":
        data["admin_note"] = inf.admin_note
    return data


@router.get("/{influencer_id}/activity")
async def activity(influencer_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """达人一站式动态:该达人的寄样 / 视频 / 投流(达人视角聚合)"""
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    coop_ids = db.scalars(select(Cooperation.id)
                          .where(Cooperation.influencer_id == inf.id)).all() or [0]

    samples = []
    sample_rows = db.execute(
        select(SampleOrder, Product).join(Product, SampleOrder.product_id == Product.id)
        .where(SampleOrder.cooperation_id.in_(coop_ids))
        .order_by(SampleOrder.created_at.desc())
    ).all()
    for row in sample_rows:
        await refresh_if_needed(db, row[0])
    for o, prod in dedupe_sample_rows(sample_rows):
        samples.append({
            "id": o.id, "product_id": prod.id, "product_name": prod.name, "status": o.status,
            "tracking_no": o.tracking_no, "courier_company": o.courier_company,
            "logistics_status": o.logistics_status,
            "signed_at": o.signed_at.isoformat() if o.signed_at else None,
            "reject_reason": o.reject_reason,
            "created_at": o.created_at.isoformat(),
        })

    videos = []
    for v, prod in db.execute(
        select(VideoTask, Product).join(Product, VideoTask.product_id == Product.id)
        .where(VideoTask.cooperation_id.in_(coop_ids))
        .order_by(VideoTask.created_at.desc())
    ).all():
        videos.append({"id": v.id, "product_id": prod.id, "product_name": prod.name, "status": v.status,
                       "blocked": v.blocked, "dy_url": v.dy_url,
                       "created_at": v.created_at.isoformat()})

    video_ids = [v["id"] for v in videos] or [0]
    promotions = []
    for p in db.scalars(select(Promotion).where(Promotion.video_task_id.in_(video_ids))
                        .order_by(Promotion.created_at.desc())).all():
        promotions.append({"id": p.id, "auth_status": p.auth_status,
                           "mode_snapshot": p.mode_snapshot, "fail_reason": p.fail_reason,
                           "created_at": p.created_at.isoformat()})

    return {"samples": samples, "videos": videos, "promotions": promotions}


# ---------- 方案B 需求1:产品合作大卡片 + 时间轴 + 全部动态 ----------

def _serialize_log(l: OperationLog) -> dict:
    return {"id": l.id, "event_type": l.event_type, "product_id": l.product_id,
            "actor_id": l.actor_id, "actor_name": l.actor_name, "actor_role": l.actor_role,
            "summary": l.summary, "detail": l.detail,
            "created_at": l.created_at.isoformat()}


@router.get("/{influencer_id}/collaborations")
def collaborations(influencer_id: int, user: User = Depends(current_user),
                   db: Session = Depends(get_db)):
    """达人 × 各产品的合作大卡片:把现有表聚合成每个产品一张卡的统计(不新建重模型)。"""
    from datetime import datetime, timedelta
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    coop_ids = db.scalars(select(Cooperation.id)
                          .where(Cooperation.influencer_id == inf.id)).all() or [0]
    # 汇总所有「有业务/被授权」的产品(授权 + 寄样 + 视频 + 出单)
    pids = set()
    pids.update(db.scalars(select(AccessGrant.product_id)
                           .where(AccessGrant.influencer_id == inf.id)).all())
    pids.update(db.scalars(select(SampleOrder.product_id)
                           .where(SampleOrder.cooperation_id.in_(coop_ids))).all())
    pids.update(db.scalars(select(VideoTask.product_id)
                           .where(VideoTask.cooperation_id.in_(coop_ids))).all())
    pids.update(db.scalars(select(OrderRecord.product_id)
                           .where(OrderRecord.influencer_id == inf.id)).all())
    pids = sorted(pid for pid in pids if pid)
    owner = db.get(User, inf.owner_bd_id) if inf.owner_bd_id else None
    if not pids:
        return {"items": [], "owner_bd_id": inf.owner_bd_id,
                "owner_bd_name": owner.display_name if owner else None,
                "commission_tier": float(inf.commission_tier)}
    products = {p.id: p for p in db.scalars(select(Product).where(Product.id.in_(pids))).all()}
    cutoff = datetime.now() - timedelta(days=30)
    cards = []
    for pid in pids:
        p = products.get(pid)
        if not p:
            continue
        grant = db.scalars(select(AccessGrant)
                           .where(AccessGrant.influencer_id == inf.id, AccessGrant.product_id == pid)
                           .order_by(AccessGrant.granted_at)).first()
        granter = db.get(User, grant.granted_by) if grant and grant.granted_by else None
        total_gmv = db.scalar(select(func.coalesce(func.sum(OrderRecord.amount), 0))
                              .where(OrderRecord.influencer_id == inf.id, OrderRecord.product_id == pid))
        gmv_30d = db.scalar(select(func.coalesce(func.sum(OrderRecord.amount), 0))
                            .where(OrderRecord.influencer_id == inf.id, OrderRecord.product_id == pid,
                                   OrderRecord.order_date >= cutoff))
        sample_total = db.scalar(select(func.count(SampleOrder.id))
                                 .where(SampleOrder.cooperation_id.in_(coop_ids),
                                        SampleOrder.product_id == pid)) or 0
        signed = db.scalar(select(func.count(SampleOrder.id))
                           .where(SampleOrder.cooperation_id.in_(coop_ids),
                                  SampleOrder.product_id == pid,
                                  SampleOrder.status == "signed")) or 0
        video_total = db.scalar(select(func.count(VideoTask.id))
                                .where(VideoTask.cooperation_id.in_(coop_ids),
                                       VideoTask.product_id == pid)) or 0
        video_pass = db.scalar(select(func.count(VideoTask.id))
                               .where(VideoTask.cooperation_id.in_(coop_ids),
                                      VideoTask.product_id == pid,
                                      VideoTask.status == "approved")) or 0
        video_fail = db.scalar(select(func.count(VideoTask.id))
                               .where(VideoTask.cooperation_id.in_(coop_ids),
                                      VideoTask.product_id == pid,
                                      VideoTask.status.in_(("rejected", "blocked")))) or 0
        promo_total = db.scalar(select(func.count(Promotion.id))
                                .join(VideoTask, Promotion.video_task_id == VideoTask.id)
                                .where(VideoTask.cooperation_id.in_(coop_ids),
                                       VideoTask.product_id == pid)) or 0
        last_promo = db.scalars(select(Promotion)
                                .join(VideoTask, Promotion.video_task_id == VideoTask.id)
                                .where(VideoTask.cooperation_id.in_(coop_ids),
                                       VideoTask.product_id == pid)
                                .order_by(Promotion.created_at.desc())).first()
        op_count = db.scalar(select(func.count(OperationLog.id))
                             .where(OperationLog.influencer_id == inf.id,
                                    OperationLog.product_id == pid)) or 0
        last_op = db.scalars(select(OperationLog)
                             .where(OperationLog.influencer_id == inf.id,
                                    OperationLog.product_id == pid)
                             .order_by(OperationLog.created_at.desc())).first()
        cards.append({
            "product_id": p.id, "product_name": p.name, "price_text": p.price_text,
            "shop_product_id": p.shop_product_id,
            "product_image": storage.thumbnail_url(p.product_image, 120) if p.product_image else None,
            "default_commission": float(p.default_commission) if p.default_commission else None,
            "influencer_commission": float(inf.commission_tier),
            "granted_by_name": granter.display_name if granter else None,
            "granted_at": grant.granted_at.isoformat() if grant else None,
            "owner_bd_name": owner.display_name if owner else None,
            "total_gmv": float(total_gmv or 0), "gmv_30d": float(gmv_30d or 0),
            "sample_total": sample_total, "sample_signed": signed,
            "video_total": video_total, "video_pass": video_pass, "video_fail": video_fail,
            "promo_total": promo_total,
            "promo_status": last_promo.auth_status if last_promo else None,
            "op_count": op_count,
            "last_op_at": last_op.created_at.isoformat() if last_op else None,
        })
    # 最近有动态的产品排前面
    cards.sort(key=lambda c: c["last_op_at"] or "", reverse=True)
    return {"items": cards, "owner_bd_id": inf.owner_bd_id,
            "owner_bd_name": owner.display_name if owner else None,
            "commission_tier": float(inf.commission_tier)}


@router.get("/{influencer_id}/collaborations/{product_id}/timeline")
def collaboration_timeline(influencer_id: int, product_id: int,
                           page: int = 1, page_size: int = 30,
                           user: User = Depends(current_user), db: Session = Depends(get_db)):
    """单个产品的合作时间轴(该达人 + 该产品的 OperationLog 倒序)。"""
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    base = select(OperationLog).where(OperationLog.influencer_id == inf.id,
                                      OperationLog.product_id == product_id)
    total = db.scalar(select(func.count()).select_from(base.subquery()))
    page = max(1, page)
    page_size = min(max(1, page_size), 100)
    rows = db.scalars(base.order_by(OperationLog.created_at.desc(), OperationLog.id.desc())
                      .offset((page - 1) * page_size).limit(page_size)).all()
    return {"items": [_serialize_log(l) for l in rows],
            "total": total, "page": page, "page_size": page_size}


@router.get("/{influencer_id}/logs")
def influencer_logs(influencer_id: int, product_id: int | None = None,
                    event_type: str | None = None, actor_id: int | None = None,
                    page: int = 1, page_size: int = 30,
                    user: User = Depends(current_user), db: Session = Depends(get_db)):
    """「全部动态」:该达人所有 OperationLog,可按产品/类型/操作人筛选 + 分页。"""
    inf = db.scalars(scope(select(Influencer).where(Influencer.id == influencer_id), user)).first()
    if not inf:
        raise HTTPException(404, "达人不存在或无权限")
    base = select(OperationLog).where(OperationLog.influencer_id == inf.id)
    if product_id is not None:
        base = base.where(OperationLog.product_id == product_id)
    if event_type:
        base = base.where(OperationLog.event_type == event_type)
    if actor_id is not None:
        base = base.where(OperationLog.actor_id == actor_id)
    total = db.scalar(select(func.count()).select_from(base.subquery()))
    page = max(1, page)
    page_size = min(max(1, page_size), 100)
    rows = db.scalars(base.order_by(OperationLog.created_at.desc(), OperationLog.id.desc())
                      .offset((page - 1) * page_size).limit(page_size)).all()
    return {"items": [_serialize_log(l) for l in rows],
            "total": total, "page": page, "page_size": page_size}
